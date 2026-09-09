"""Build the sample decks used by the test suite.

Run with ``python tests/fixtures/make_fixtures.py``. The generated files are small
synthetic pitch decks for a fictional company; they exist so the ingestion tests
exercise real PDF and PPTX parsing rather than mocks.
"""

from __future__ import annotations

from pathlib import Path

FIXTURE_DIR = Path(__file__).resolve().parent

#: Slide content shared by both fixture formats, as (title, body lines).
DECK_SLIDES = [
    (
        "Northwind Diagnostics",
        [
            "AI-native pathology for community hospitals",
            "Series A round | San Diego, CA | Founded 2021",
            "www.northwinddx.com",
        ],
    ),
    (
        "The Problem",
        [
            "Community hospitals wait 9 days for pathology results.",
            "Diagnostic delay drives $4.2B in avoidable annual cost.",
            "62% of rural hospitals have no on-site pathologist.",
        ],
    ),
    (
        "Our Solution",
        [
            "FDA-cleared digital pathology workflow with automated triage.",
            "Deployed as SaaS on existing scanner hardware.",
            "Gross margin of 78% at current volume.",
        ],
    ),
    (
        "Traction",
        [
            "ARR of $4.2M, up from $1.6M the prior year.",
            "Revenue growth of 163% year over year.",
            "42 hospital customers under multi-year contracts.",
            "Net revenue retention of 118%.",
            "Recurring revenue is 92% of total revenue.",
        ],
    ),
    (
        "Market",
        [
            "TAM of $18B across US diagnostic pathology.",
            "SAM of $3.2B in community and rural hospital segments.",
            "Market growing 14% annually.",
        ],
    ),
    (
        "Technology and IP",
        [
            "7 issued patents and 4 patent applications pending.",
            "FDA 510(k) clearance granted for the triage module.",
            "Proprietary dataset of 2.1M annotated slides.",
        ],
    ),
    (
        "Financials",
        [
            "Monthly burn of $310,000.",
            "Cash on hand of $5.4M.",
            "Runway of 17 months.",
            "Capital raised to date: $11.5M.",
            "Gross margin 78%. EBITDA margin -45%.",
        ],
    ),
    (
        "The Ask",
        [
            "Raising $12M Series A.",
            "Top customer represents 11% of revenue.",
            "Use of proceeds: commercial expansion and regulatory submissions.",
        ],
    ),
]

FINANCIAL_TABLE = [
    ["Metric", "FY2023", "FY2024"],
    ["Revenue", "$1.6M", "$4.2M"],
    ["Gross margin", "71%", "78%"],
    ["Customers", "18", "42"],
]


def build_pdf(path: Path) -> Path:
    """A text-layer PDF deck, one slide per page, with a financial table."""
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas

    page_size = landscape(letter)
    pdf = canvas.Canvas(str(path), pagesize=page_size)
    width, height = page_size

    for title, lines in DECK_SLIDES:
        pdf.setFont("Helvetica-Bold", 26)
        pdf.drawString(inch, height - inch, title)
        pdf.setFont("Helvetica", 15)
        y = height - 1.7 * inch
        for line in lines:
            pdf.drawString(inch, y, line)
            y -= 0.42 * inch
        pdf.showPage()

    # A dedicated table page so table extraction has something real to find.
    pdf.setFont("Helvetica-Bold", 26)
    pdf.drawString(inch, height - inch, "Financial Summary")
    y = height - 1.8 * inch
    for row in FINANCIAL_TABLE:
        x = inch
        for cell in row:
            pdf.setFont("Helvetica", 14)
            pdf.drawString(x, y, cell)
            x += 2.4 * inch
        y -= 0.4 * inch
    pdf.showPage()
    pdf.save()
    return path


def build_pptx(path: Path) -> Path:
    """A PPTX deck with titles, bodies, speaker notes and a real table shape."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]
    title_body_layout = presentation.slide_layouts[1]

    for title, lines in DECK_SLIDES:
        slide = presentation.slides.add_slide(title_body_layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = lines[0]
        for line in lines[1:]:
            body.add_paragraph().text = line
        slide.notes_slide.notes_text_frame.text = (
            f"Speaker notes for {title}: figures are as of the most recent quarter."
        )

    # Table slide.
    slide = presentation.slides.add_slide(blank_layout)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(6), Inches(0.6))
    box.text_frame.text = "Financial Summary"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

    rows, cols = len(FINANCIAL_TABLE), len(FINANCIAL_TABLE[0])
    table = slide.shapes.add_table(
        rows, cols, Inches(0.5), Inches(1.2), Inches(8), Inches(2.4)
    ).table
    for r, row in enumerate(FINANCIAL_TABLE):
        for c, cell in enumerate(row):
            table.cell(r, c).text = cell

    presentation.save(str(path))
    return path


def build_image_only_pdf(path: Path) -> Path:
    """A PDF with no text layer at all, for the image-only error path."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setFillColorRGB(0.85, 0.85, 0.85)
    pdf.rect(100, 300, 400, 300, fill=1, stroke=0)
    pdf.showPage()
    pdf.save()
    return path


def build_corrupt_pdf(path: Path) -> Path:
    """A file with a PDF extension that is not a PDF."""
    path.write_bytes(b"%PDF-1.4\nthis is not actually a valid pdf body\n")
    return path


def build_legacy_ppt(path: Path) -> Path:
    """A file carrying the OLE signature of a legacy binary .ppt."""
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    return path


def build_comparables_csv(path: Path) -> Path:
    path.write_text(
        "Company,Buyer,Buyer Type,Date,Deal Value,Revenue,EBITDA,Growth,Rationale\n"
        "PathAI Systems,Roche,Strategic,2023-06-01,$320M,$40M,$4M,55%,Diagnostics platform fit\n"
        "SlideScan,Philips,Strategic,2022-11-15,$185M,$22M,,48%,Scanner ecosystem\n"
        "HistoCloud,Thermo Fisher,Strategic,2024-02-20,$450M,$50M,$8M,62%,Lab network expansion\n"
        "DiagnostiQ,Vista Equity,Financial,2023-09-10,$210M,$35M,$7M,25%,Recurring revenue roll-up\n"
        "CellSight,Sponsor Partners,Financial,2022-04-01,$95M,$19M,$3M,18%,Buy and build\n",
        encoding="utf-8",
    )
    return path


def main() -> None:
    outputs = [
        build_pdf(FIXTURE_DIR / "sample_deck.pdf"),
        build_pptx(FIXTURE_DIR / "sample_deck.pptx"),
        build_image_only_pdf(FIXTURE_DIR / "image_only.pdf"),
        build_corrupt_pdf(FIXTURE_DIR / "corrupt.pdf"),
        build_legacy_ppt(FIXTURE_DIR / "legacy.ppt"),
        build_comparables_csv(FIXTURE_DIR / "comparables.csv"),
    ]
    for path in outputs:
        print(f"wrote {path.name} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
