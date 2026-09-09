"""PowerPoint ingestion.

python-pptx reads the Open XML format (.pptx) natively, including grouped shapes,
tables and speaker notes. The legacy binary .ppt format is not Open XML and cannot
be read by python-pptx; that case is detected up front and reported with a clear
remedy rather than failing deep inside the parser.
"""

from __future__ import annotations

import logging
import re
import zipfile
from pathlib import Path
from typing import List, Optional

from .base import (
    DocumentExtractionError,
    ExtractedDocument,
    ExtractedPage,
    ExtractedTable,
    ExtractionMethod,
    SourceKind,
    UnsupportedFormatError,
)

logger = logging.getLogger(__name__)

#: OLE compound-document signature: the marker of a legacy binary .ppt file.
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _clean_text(text: Optional[str]) -> str:
    if not text:
        return ""
    lines = []
    for line in text.replace("\x0b", "\n").splitlines():
        collapsed = re.sub(r"[ \t ]+", " ", line).strip()
        if collapsed:
            lines.append(collapsed)
    return "\n".join(lines)


def _iter_shapes(shapes):
    """Walk shapes recursively so text inside groups is not missed."""
    for shape in shapes:
        # Group shapes (type 6) nest their own shape tree.
        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _table_from_shape(shape, slide_number: int) -> Optional[ExtractedTable]:
    try:
        rows = [
            [_clean_text(cell.text).replace("\n", " ") for cell in row.cells]
            for row in shape.table.rows
        ]
    except Exception as exc:
        logger.debug("Table read failed on slide %s: %s", slide_number, exc)
        return None
    table = ExtractedTable(page_or_slide=slide_number, rows=rows)
    return table if table.is_meaningful else None


def _slide_title(slide, body_text: str) -> Optional[str]:
    """Prefer the slide's declared title placeholder; fall back to the first line."""
    try:
        if slide.shapes.title is not None:
            title = _clean_text(slide.shapes.title.text)
            if title:
                return title.splitlines()[0][:90]
    except (AttributeError, ValueError):
        pass
    for line in body_text.splitlines():
        stripped = line.strip()
        if 3 <= len(stripped) <= 90:
            return stripped
    return None


def _speaker_notes(slide, slide_number: int, doc: ExtractedDocument) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
        return _clean_text(frame.text) if frame is not None else ""
    except Exception as exc:
        logger.debug("Notes read failed on slide %s: %s", slide_number, exc)
        doc.warnings.append(f"Slide {slide_number}: speaker notes could not be read.")
        return ""


def _detect_legacy_ppt(path: Path) -> bool:
    try:
        with path.open("rb") as fh:
            return fh.read(8) == _OLE_MAGIC
    except OSError:
        return False


def parse_pptx(path: str | Path) -> ExtractedDocument:
    """Extract slide text, tables, titles and speaker notes from a PowerPoint deck."""
    path = Path(path)
    if not path.exists():
        raise DocumentExtractionError(f"{path.name} was not found.", "Upload the file again.")
    if path.stat().st_size == 0:
        raise DocumentExtractionError(
            f"{path.name} is empty (0 bytes).", "Upload a complete PowerPoint file."
        )

    if _detect_legacy_ppt(path):
        raise UnsupportedFormatError(
            path.name,
            "It is a legacy binary PowerPoint (.ppt) file, which cannot be read directly.",
        )

    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    try:
        presentation = Presentation(str(path))
    except PackageNotFoundError as exc:
        raise UnsupportedFormatError(
            path.name, "It is not a valid PowerPoint package."
        ) from exc
    except zipfile.BadZipFile as exc:
        raise DocumentExtractionError(
            f"{path.name} appears to be corrupted.",
            "Re-save the deck from PowerPoint and upload it again.",
        ) from exc
    except Exception as exc:
        raise DocumentExtractionError(
            f"{path.name} could not be opened ({exc}).",
            "Re-save the deck as .pptx and upload it again.",
        ) from exc

    doc = ExtractedDocument(filename=path.name, kind=SourceKind.PPTX)

    slides = list(presentation.slides)
    if not slides:
        raise DocumentExtractionError(
            f"{path.name} contains no slides.", "Upload a deck with at least one slide."
        )

    for index, slide in enumerate(slides, start=1):
        fragments: List[str] = []
        tables: List[ExtractedTable] = []

        for shape in _iter_shapes(slide.shapes):
            try:
                if getattr(shape, "has_table", False):
                    table = _table_from_shape(shape, index)
                    if table is not None:
                        tables.append(table)
                    continue
                if getattr(shape, "has_text_frame", False):
                    text = _clean_text(shape.text_frame.text)
                    if text:
                        fragments.append(text)
                # Charts carry their category names, which often hold the metric labels.
                if getattr(shape, "has_chart", False):
                    try:
                        categories = [
                            str(c) for c in shape.chart.plots[0].categories if c is not None
                        ]
                        if categories:
                            fragments.append(" ".join(categories))
                    except Exception:
                        logger.debug("Chart categories unavailable on slide %s", index)
            except Exception as exc:
                logger.debug("Shape read failed on slide %s: %s", index, exc)
                continue

        body = "\n".join(fragments)
        notes = _speaker_notes(slide, index, doc)

        if body:
            method = ExtractionMethod.TEXT_LAYER
        elif tables:
            method = ExtractionMethod.TABLE
        elif notes:
            method = ExtractionMethod.SPEAKER_NOTES
        else:
            method = ExtractionMethod.NONE

        doc.pages.append(
            ExtractedPage(
                number=index,
                title=_slide_title(slide, body),
                text=body,
                notes=notes,
                tables=tables,
                method=method,
            )
        )

    image_only = [p.number for p in doc.pages if not p.searchable_text()]
    if image_only:
        doc.warnings.append(
            f"{len(image_only)} slide(s) contain only images with no readable text: "
            f"slide(s) {', '.join(str(n) for n in image_only[:12])}"
            + ("..." if len(image_only) > 12 else "")
            + ". Enter any values from those slides manually."
        )

    if doc.is_empty:
        raise DocumentExtractionError(
            f"No text could be extracted from {path.name}.",
            "Every slide appears to be an image. Enter the company details manually "
            "in the Company Inputs tab.",
        )
    return doc


def parse_document(path: str | Path, allow_ocr: bool = True) -> ExtractedDocument:
    """Dispatch to the right parser based on file extension.

    This is the single entry point the application uses for any uploaded deck.
    """
    from .pdf_parser import parse_pdf

    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path, allow_ocr=allow_ocr)
    if suffix in {".pptx", ".pptm"}:
        return parse_pptx(path)
    if suffix == ".ppt":
        raise UnsupportedFormatError(
            path.name,
            "Legacy .ppt files cannot be read directly. Open the deck in PowerPoint "
            "and save it as .pptx or export it as a PDF.",
        )
    raise UnsupportedFormatError(path.name, f"Unrecognised extension '{suffix}'.")
